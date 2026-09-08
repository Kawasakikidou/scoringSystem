#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成二期 Excel 验收样例 samples/名单样例-四维.xlsx 与 .xls：
内容与 samples/名单样例-乱序.txt 逐行对齐（乱序多列 + 坏行 + 重复行），
统计口径应与 txt 完全一致：成功 10 行（新增 9、更新 1）、跳过 3 行。

约定（覆盖已知坑）：
- 恰好 10 位数字的单元格写为“数值单元格”（学号若被 Excel 当数字，读回禁止科学计数法）；
- 其余一律写文本（避免 18 位连写数字等被浮点精度破坏）；
- 空行不写；表头/题头行保留为文本行。
依赖：python3 + openpyxl（xlsx）+ xlwt（xls）；缺失时 pip install openpyxl xlwt
"""
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "samples", "名单样例-乱序.txt")
OUT_XLSX = os.path.join(ROOT, "samples", "名单样例-四维.xlsx")
OUT_XLS = os.path.join(ROOT, "samples", "名单样例-四维.xls")

def load_rows():
    """每行 = 该行文本按空白/逗号/顿号切出的单元格列表；空行跳过。"""
    rows = []
    with open(SRC, encoding="utf-8") as fh:
        for raw in fh.read().splitlines():
            line = raw.strip()
            if not line:
                continue
            tokens = [t for t in re.split(r"[\s,、]+", line) if t]
            rows.append(tokens)
    return rows

def is_numeric_no(tok):
    return bool(re.fullmatch(r"\d{10}", tok))

def gen_xlsx(rows):
    try:
        from openpyxl import Workbook
    except ImportError:
        sys.exit("缺少 openpyxl：pip install openpyxl")
    wb = Workbook()
    ws = wb.active
    ws.title = "名单"
    for r, toks in enumerate(rows, start=1):
        for c, tok in enumerate(toks, start=1):
            ws.cell(row=r, column=c, value=int(tok) if is_numeric_no(tok) else tok)
    wb.save(OUT_XLSX)
    print("已生成", OUT_XLSX)

def gen_xls(rows):
    try:
        import xlwt
    except ImportError:
        sys.exit("缺少 xlwt：pip install xlwt")
    wb = xlwt.Workbook(encoding="utf-8")
    sh = wb.add_sheet("名单")
    for r, toks in enumerate(rows):
        for c, tok in enumerate(toks):
            sh.write(r, c, int(tok) if is_numeric_no(tok) else tok)
    wb.save(OUT_XLS)
    print("已生成", OUT_XLS)

if __name__ == "__main__":
    rows = load_rows()
    print("源文件行（非空）:", len(rows))
    gen_xlsx(rows)
    gen_xls(rows)
    print("完成：两个 Excel 样例已按 txt 口径生成，导入统计应与 txt 一致（成功10/新增9/更新1/跳过3）。")
